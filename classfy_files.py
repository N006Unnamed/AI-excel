import os
import json
import pandas as pd
import numpy as np
import matplotlib.pyplot as plt
import openpyxl
from openai import OpenAI
from datetime import datetime
import re


class DeepSeekCostProcessor:
    def __init__(self, api_key):
        self.client = OpenAI(
            api_key=api_key,
            base_url="https://api.deepseek.com/v1",
        )
        self.cost_details = None
        self.cost_template = None
        self.category_mapping = None
        self.years = [str(y) for y in range(2025, 2046)]  # 2025-2045

    def load_excel_data(self, cost_details_path, template_path):
        """从Excel文件加载费用明细和分类模板"""
        try:
            # 加载费用明细表
            self.cost_details = pd.read_excel(cost_details_path)
            print(f"成功加载费用明细表: {cost_details_path}")
            print(f"包含 {len(self.cost_details)} 条记录")

            # 加载分类模板
            self.cost_template = pd.read_excel(template_path)
            print(f"成功加载分类模板: {template_path}")
            print(f"包含 {len(self.cost_template)} 条记录")

            # 提取年份列
            year_cols = [col for col in self.cost_details.columns if re.match(r'^\d{4}$', str(col))]
            if year_cols:
                self.years = sorted(year_cols, key=lambda x: int(x))
                print(f"检测到年份列: {self.years}")

            return True
        except Exception as e:
            print(f"Excel加载失败: {e}")
            return False

    def extract_data_for_prompt(self):
        """从DataFrame提取数据构建API提示词"""
        # 提取费用明细
        cost_details_text = "费用明细表:\n"
        for _, row in self.cost_details.iterrows():
            # 获取序号和科目
            seq = row.get('序号', '')
            category = row.get('科目', '')
            cost_details_text += f"{seq}. {category}\n"

        # 提取分类模板
        template_text = "分类模板:\n"
        for _, row in self.cost_template.iterrows():
            # 获取序号和科目
            seq = row.get('序号', '')
            category = row.get('科目', '')

            # 处理子类别的特殊格式
            if isinstance(seq, str) and '.' in seq:
                template_text += f"{seq} {category}\n"
            else:
                template_text += f"{seq}. {category}\n"

        return cost_details_text, template_text

    def get_cost_mapping(self):
        """使用DeepSeek API获取费用分类映射"""
        # 提取数据构建提示词
        cost_details_text, template_text = self.extract_data_for_prompt()

        # 构建提示词
        prompt = f"""
您是一位财务分析专家，请将费用明细表中的项目按分类模板归类，并返回JSON格式的映射关系。

{cost_details_text}

{template_text}

输出要求：
1. 返回严格JSON格式：{{"分类名称": [对应明细序号列表]}}
2. 示例：{{"外购燃料及动力费": [5,6,7]}}
3. 注意：其他费用需指定子类(5.1/5.2/5.3)
4. 不要包含任何额外解释或文本
5. 确保所有费用明细项目都被分类
"""

        # 调用DeepSeek API
        try:
            response = self.client.chat.completions.create(
                model="deepseek-chat",
                messages=[{"role": "user", "content": prompt}],
                temperature=0.3,
                max_tokens=2000
            )
            content = response.choices[0].message.content
            print("API响应内容:")
            print(content[:500] + "..." if len(content) > 500 else content)

            # 尝试解析JSON
            try:
                mapping = json.loads(content)
                self.validate_mapping(mapping)
                self.category_mapping = mapping
                return mapping
            except json.JSONDecodeError:
                # 如果返回内容包含额外文本，尝试提取JSON
                try:
                    start = content.find("{")
                    end = content.rfind("}") + 1
                    mapping = json.loads(content[start:end])
                    self.validate_mapping(mapping)
                    self.category_mapping = mapping
                    return mapping
                except Exception as e:
                    print(f"JSON解析失败: {e}")
                    return self.get_fallback_mapping()
        except Exception as e:
            print(f"API调用失败: {e}")
            return self.get_fallback_mapping()

    def validate_mapping(self, mapping):
        """验证映射关系完整性"""
        print("验证分类映射...")

        # 获取所有费用明细序号
        all_detail_ids = set(self.cost_details['序号'].astype(int).unique())
        classified_ids = set()

        # 检查每个分类
        for category, ids in mapping.items():
            if not isinstance(ids, list):
                raise ValueError(f"分类 '{category}' 的值不是列表")

            for id_val in ids:
                if not isinstance(id_val, int):
                    raise ValueError(f"分类 '{category}' 包含非整数值: {id_val}")
                classified_ids.add(id_val)

        # 检查未分类的项目
        unclassified = all_detail_ids - classified_ids
        if unclassified:
            print(f"警告: 以下序号未被分类: {sorted(unclassified)}")

        # 检查不存在的序号
        invalid_ids = classified_ids - all_detail_ids
        if invalid_ids:
            print(f"警告: 以下序号在费用明细中不存在: {sorted(invalid_ids)}")

        print("分类映射验证完成")

    def get_fallback_mapping(self):
        """备用映射关系（当API失败时使用）"""
        print("使用备用分类映射")
        return {
            "外购燃料及动力费": [5, 6, 7],
            "工资福利费": [13],
            "修理费": [2, 3, 4, 9, 10, 11, 17],
            "其他管理费用": [1, 14, 15, 16, 18, 19, 20],
            "其他营业费用": [8, 12]
        }

    def process_and_fill_data(self, output_path):
        """处理数据并填充模板"""
        if self.category_mapping is None:
            self.get_cost_mapping()

        # 创建结果DataFrame的副本
        result_df = self.cost_template.copy()

        # 初始化汇总数据
        category_totals = {}
        for category in self.category_mapping:
            category_totals[category] = pd.Series(0, index=self.years)

        # 按类别汇总数据
        for category, ids in self.category_mapping.items():
            # 筛选当前类别的行
            category_rows = self.cost_details[self.cost_details['序号'].isin(ids)]

            if not category_rows.empty:
                # 按年求和
                yearly_sum = category_rows[self.years].sum()
                category_totals[category] = yearly_sum
                print(f"分类 '{category}' 汇总完成: {len(ids)} 个项目")

        # 特殊规则处理：其他营业费用在建设期（2025）为0
        if "其他营业费用" in category_totals:
            category_totals["其他营业费用"]["2025"] = 0
            print("应用特殊规则: 2025年其他营业费用=0")

        # 更新分类模板数据
        for index, row in result_df.iterrows():
            category = row['科目']

            if category in category_totals:
                for year in self.years:
                    result_df.at[index, year] = category_totals[category][year]

        # 计算衍生字段
        self.calculate_derived_fields(result_df)

        # 保存结果
        result_df.to_excel(output_path, index=False)
        print(f"结果已保存至: {output_path}")

        # 生成可视化
        self.generate_visualization(result_df, output_path)

        return result_df

    def calculate_derived_fields(self, result_df):
        """计算经营成本、总成本等衍生字段"""
        print("计算衍生字段...")

        # 获取各主要类别的行索引
        category_indices = {}
        for category in ["外购原材料费", "外购燃料及动力费", "工资福利费", "修理费", "其他费用",
                         "经营成本", "总成本费用", "折旧费", "摊销费", "利息支出"]:
            matches = result_df[result_df['科目'] == category]
            if not matches.empty:
                category_indices[category] = matches.index[0]

        for year in self.years:
            # 计算其他费用 = 5.1 + 5.2 + 5.3
            other_manufacturing = result_df.loc[result_df['科目'] == "其他制造费用", year].values
            other_management = result_df.loc[result_df['科目'] == "其他管理费用", year].values
            other_operating = result_df.loc[result_df['科目'] == "其他营业费用", year].values

            if other_manufacturing.size and other_management.size and other_operating.size:
                other_total = other_manufacturing[0] + other_management[0] + other_operating[0]
                if "其他费用" in category_indices:
                    result_df.at[category_indices["其他费用"], year] = other_total

            # 计算经营成本 = 1+2+3+4+5
            if all(cat in category_indices for cat in
                   ["外购原材料费", "外购燃料及动力费", "工资福利费", "修理费", "其他费用"]):
                cost_1 = result_df.at[category_indices["外购原材料费"], year]
                cost_2 = result_df.at[category_indices["外购燃料及动力费"], year]
                cost_3 = result_df.at[category_indices["工资福利费"], year]
                cost_4 = result_df.at[category_indices["修理费"], year]
                cost_5 = result_df.at[category_indices["其他费用"], year]
                if "经营成本" in category_indices:
                    result_df.at[category_indices["经营成本"], year] = cost_1 + cost_2 + cost_3 + cost_4 + cost_5

            # 计算总成本费用 = 经营成本 + 折旧费 + 摊销费 + 利息支出
            if all(cat in category_indices for cat in ["经营成本", "折旧费", "摊销费", "利息支出"]):
                operating_cost = result_df.at[category_indices["经营成本"], year]
                depreciation = result_df.at[category_indices["折旧费"], year]
                amortization = result_df.at[category_indices["摊销费"], year]
                interest = result_df.at[category_indices["利息支出"], year]
                if "总成本费用" in category_indices:
                    result_df.at[
                        category_indices["总成本费用"], year] = operating_cost + depreciation + amortization + interest

        print("衍生字段计算完成")

    def generate_visualization(self, result_df, output_path):
        """生成成本趋势可视化图表"""
        print("生成可视化图表...")

        try:
            # 提取主要成本类别
            main_categories = [
                "外购燃料及动力费", "工资福利费", "修理费",
                "其他管理费用", "其他营业费用"
            ]

            # 准备图表数据
            chart_data = {}
            for category in main_categories:
                category_row = result_df[result_df['科目'] == category]
                if not category_row.empty:
                    values = category_row[self.years].values.flatten()
                    chart_data[category] = values

            if not chart_data:
                print("无有效数据生成图表")
                return

            # 创建堆叠面积图
            plt.figure(figsize=(14, 8))

            # 准备堆叠数据
            labels = list(chart_data.keys())
            data_values = list(chart_data.values())

            # 绘制堆叠面积图
            plt.stackplot(self.years, *data_values, labels=labels)
            plt.title("2025-2045年成本结构趋势分析", fontsize=15)
            plt.xlabel("年份", fontsize=12)
            plt.ylabel("成本金额", fontsize=12)
            plt.legend(loc="upper left")
            plt.xticks(rotation=45)
            plt.tight_layout()

            # 保存图表
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            chart_file = f"cost_trend_{timestamp}.png"
            plt.savefig(chart_file)
            print(f"成本趋势分析图已保存: {chart_file}")

            # 添加图表到Excel
            self.add_chart_to_excel(output_path, chart_file)
        except Exception as e:
            print(f"图表生成失败: {e}")

    def add_chart_to_excel(self, excel_path, chart_path):
        """将图表添加到Excel文件"""
        try:
            # 加载工作簿
            wb = load_work
            book(excel_path)
            ws = wb.active

            # 创建图像对象
            img = openpyxl.drawing.image.Image(chart_path)

            # 调整图像大小
            img.width = 800
            img.height = 500

            # 确定插入位置（最后一行之后）
            max_row = ws.max_row
            anchor = f"A{max_row + 3}"

            # 添加图像到工作表
            ws.add_image(img, anchor)

            # 保存工作簿
            wb.save(excel_path)
            print(f"图表已添加到Excel文件: {excel_path}")
        except Exception as e:
            print(f"添加图表到Excel失败: {e}")


# 主函数
def main():
    # 配置参数
    API_KEY = os.getenv("DEEPSEEK_API_KEY")
    if not API_KEY:
        print("错误: 未设置DEEPSEEK_API_KEY环境变量")
        return

    # 文件路径 - 替换为您的实际文件路径
    COST_DETAILS_PATH = "费用明细表.xlsx"
    TEMPLATE_PATH = "分类模板.xlsx"
    OUTPUT_PATH = f"成本汇总结果_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"

    # 创建处理器
    processor = DeepSeekCostProcessor(API_KEY)

    # 加载Excel数据
    if not processor.load_excel_data(COST_DETAILS_PATH, TEMPLATE_PATH):
        print("数据加载失败，程序终止")
        return

    # 处理数据并生成结果
    result_df = processor.process_and_fill_data(OUTPUT_PATH)

    # 打印部分结果
    print("\n处理完成! 结果预览:")
    print(result_df[['序号', '科目'] + processor.years[:3]].head(10))


if __name__ == "__main__":
    main()